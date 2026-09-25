#!/usr/bin/env python3
"""
Verification harness for the T1 Python-I question paper (guidelines
PYTHON-1_T1_QUESTION_PAPER_GENERATION_GUIDELINES.pdf).

1. RUNTIME   every output MCQ is executed and the marked option compared with
             the real output; the descriptive questions' stated examples are
             executed; a reference solution of the online question is run to
             prove the specification is well posed.
2. RULES     §3 unit blueprint, §4/§7 set rules, §5 PB split, §6 six options
             with e) Error and f) None of the above, §9 one integrated online
             question, §10 template controls, §11 confidentiality.
"""

import contextlib
import io
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_t1_qp import (MCQS, OFFLINE_UNIT_TARGET, OUTDIR, Q2, Q3,  # noqa: E402
                         SET_ORDER, T1_UNIT_BLUEPRINT)

FAIL = []

RUNTIME = {54: "6", 107: "4", 165: "12", 167: "2 5", 176: "1 5", 160: "2 5 7"}
EXPRESSIONS = {64: [512, 64, 512]}
ERRORS = {86: "ZeroDivisionError"}
MANUAL = {8, 21}          # concept question and paper algorithm, traced by hand


def check(cond, msg):
    print(("  OK   " if cond else "  FAIL ") + msg)
    if not cond:
        FAIL.append(msg)


def run(src):
    buf = io.StringIO()
    try:
        with contextlib.redirect_stdout(buf):
            exec(src, {})
        return buf.getvalue(), None
    except Exception as exc:                       # noqa: BLE001
        return buf.getvalue(), type(exc).__name__


def numbers(s):
    return [int(n) for n in re.findall(r"-?\d+", s)]


# --------------------------------------------------------------------------- #
def verify_answers():
    print("== 1. every MCQ answer re-derived from the code itself ==")
    covered = set(RUNTIME) | set(EXPRESSIONS) | set(ERRORS) | MANUAL
    pbs = [m["pb"] for m in MCQS]
    check(set(pbs) == covered and len(pbs) == len(set(pbs)),
          f"all {len(pbs)} MCQs covered by exactly one proof method")
    for m in MCQS:
        pb, correct = m["pb"], m["opts"][m["ans"]]
        check(len(m["opts"]) == 6 and len(set(m["opts"])) == 6,
              f"PB{pb}: exactly 6 distinct options (§6)")
        if m["is_output"]:
            check(m["opts"][4] == "Error" and m["opts"][5] == "None of the above",
                  f"PB{pb}: e) Error and f) None of the above (§6)")
            check(m["ans"] < 5 or pb == 86,
                  f"PB{pb}: correct answer is a real option, not f)")
        if pb in MANUAL:
            print(f"  ---- PB{pb}: hand traced, marked {correct!r}")
            continue
        if pb in EXPRESSIONS:
            vals = [eval(line, {}) for line in m["code"]]     # noqa: S307
            check(vals == EXPRESSIONS[pb] == [int(x) for x in correct.split(",")],
                  f"PB{pb}: expressions = {vals}, marked {correct!r}")
            continue
        out, err = run("\n".join(m["code"]))
        if pb in ERRORS:
            check(err == ERRORS[pb] and correct == "Error",
                  f"PB{pb}: raises {err}, marked {correct!r}")
            continue
        check(out.strip().rstrip(",") == RUNTIME[pb] and
              numbers(out) == numbers(correct),
              f"PB{pb}: printed {out.strip()!r}, marked {correct!r}")

    print("\n== descriptive / online questions are well posed ==")
    out, err = run("total = 7 / 2\ncount = int(total)\n"
                   "result = count ** 2 + 9 % 4\n"
                   "print(result, type(result), type(total))")
    check(out.strip() == "10 <class 'int'> <class 'float'>",
          f"Q-2 1) cell really prints {out.strip()!r}")
    n = 9875
    while n >= 10:
        s = 0
        while n:
            s += n % 10
            n //= 10
        n = s
    check(n == 2, "Q-2 2) digital root of 9875 is 2, as the paper states")

    def bill(prev, cur, kind):
        units = cur - prev
        if kind == 1:
            if units <= 100:
                energy = units * 3
            elif units <= 200:
                energy = 300 + (units - 100) * 5
            else:
                energy = 300 + 500 + (units - 200) * 7
            service = 50
        else:
            energy, service = units * 9, 150
        return units, energy, service, energy + service
    cases = [(1000, 1450, 1), (1000, 1450, 2), (500, 650, 1), (500, 500, 1)]
    for prev, cur, kind in cases:
        u, e, s, t = bill(prev, cur, kind)
        assert t == e + s
    check(bill(1000, 1450, 1) == (450, 2550, 50, 2600),
          "Q-3 reference solution: domestic 450 units -> bill 2600 (HIGH USAGE)")
    check(bill(1000, 1450, 2) == (450, 4050, 150, 4200),
          "Q-3 reference solution: commercial 450 units -> bill 4200")
    check(bill(500, 650, 1) == (150, 550, 50, 600),
          "Q-3 reference solution: 150 units straddles the 100-unit slab")


# --------------------------------------------------------------------------- #
def set_box_letters(doc):
    return ["".join(t.text or "" for t in tb.iter(qn("w:t"))).strip()
            for tb in doc.element.body.iter(qn("w:txbxContent"))
            if "".join(t.text or "" for t in tb.iter(qn("w:t"))).strip()
            .startswith("Set")]


def parse_options(cell_text):
    return dict(re.findall(r"\(([a-f])\)\s*([^\t\n]+)", cell_text))


def verify_documents():
    print("\n== 2. the generated student papers ==")
    audit = (OUTDIR / "AUDIT SHEET AND ANSWER KEY_T1_PYTHON-I_MDP_faculty only.txt")
    audit_txt = audit.read_text()
    by_opts = {frozenset(m["opts"]): m for m in MCQS}
    check(len(by_opts) == len(MCQS), "option sets identify each MCQ uniquely")

    unit_offline = {1: 0, 2: 0, 3: 0}
    for m in MCQS:
        unit_offline[m["unit"]] += 1
    for q in Q2:
        unit_offline[q["unit"]] += q["marks"]

    for letter in ("A", "B", "C"):
        path = OUTDIR / f"SET {letter}_T1_PYTHON-1_TEST PAPER_SEM I_MDP.docx"
        doc = Document(str(path))
        print(f"-- SET {letter} ({path.name})")
        boxes = set_box_letters(doc)
        check(len(boxes) == 4 and all(b == f"Set: {letter}" for b in boxes),
              f"all four 'Set' boxes read 'Set: {letter}'")

        head = " ".join(c.text for r in doc.tables[0].rows for c in r.cells)
        for probe in ("TEST 1 ( CO1 ) (OFFLINE)", "MAX MARKS: 16 marks",
                      "2:15 to 3:30 pm", "29-Sep-2026", "1.25", "017012194",
                      "017212194", "CE/IT/CSD", "Computer Programming using "
                      "Python-I"):
            check(probe in head, f"offline header contains {probe!r}")
        check("2/3" not in head and "CO2" not in head,
              "no leftover 2/3 or CO2 in the offline heading")
        head_on = " ".join(c.text for r in doc.tables[3].rows for c in r.cells)
        for probe in ("TEST 1 ( CO1 ) (ONLINE)", "MAX MARKS: 09 marks",
                      "4:15 to 5:30 pm"):
            check(probe in head_on, f"online header contains {probe!r}")

        # ---- offline table -------------------------------------------------
        rows = doc.tables[1].rows
        block = rows[1]
        check(block.cells[0].text.strip() == "Q-1" and
              "1 Mark each" in block.cells[2].text and
              block.cells[3].text.strip() == "(10)",
              f"Q-1 block row: compulsory MCQ, 1 mark each, marks column "
              f"{block.cells[3].text.strip()!r}")
        mcq_rows, desc = [], 0
        for r in rows[2:]:
            sub, marks, body = (r.cells[1].text.strip(),
                                r.cells[3].text.strip(), r.cells[2].text)
            if re.search(r"\([a-f]\)", body):
                mcq_rows.append((sub, body, marks))
            else:
                desc += int(marks)
        check(len(mcq_rows) == 10 and sum(int(m) for _, _, m in mcq_rows) == 10,
              f"{len(mcq_rows)} MCQs worth "
              f"{sum(int(m) for _, _, m in mcq_rows)} marks from the PB (§5)")
        check(desc == 6, f"outside-PB descriptive = {desc} marks (§5)")
        check(all(len(parse_options(b)) == 6 for _, b, _ in mcq_rows),
              "every MCQ prints six options (§6)")
        check([by_opts[frozenset(parse_options(b).values())]["pb"]
               for _, b, _ in mcq_rows]
              == [MCQS[i]["pb"] for i in SET_ORDER[letter]],
              "MCQ sequence matches SET_ORDER; nothing else reordered")
        for sub, body, _ in mcq_rows:
            m = by_opts[frozenset(parse_options(body).values())]
            if m["is_output"]:
                o = parse_options(body)
                check(o["e"] == "Error" and o["f"] == "None of the above",
                      f"{sub} (PB{m['pb']}) prints e) Error and "
                      f"f) None of the above")
        # monospaced code with preserved indentation (§10)
        nested = next(r for r in rows
                      if "elif a > 5:" in r.cells[2].text)
        code_paras, text_paras = [], []
        for para in nested.cells[2].paragraphs:
            pPr = para._p.find(qn("w:pPr"))
            ind = pPr.find(qn("w:ind")) if pPr is not None else None
            fonts = {r._r.find(qn("w:rPr")).find(qn("w:rFonts")).get(qn("w:ascii"))
                     for r in para.runs if r._r.find(qn("w:rPr")) is not None
                     and r._r.find(qn("w:rPr")).find(qn("w:rFonts")) is not None}
            (code_paras if ind is not None else text_paras).append(
                (int(ind.get(qn("w:left"))) if ind is not None else None, fonts))
        check(code_paras and all(f == {"Courier New"} for _, f in code_paras),
              f"all {len(code_paras)} code lines are monospaced (§10)")
        check(all(f == {"Times New Roman"} for _, f in text_paras),
              "stem and options keep the template font")
        levels = sorted({i for i, _ in code_paras})
        check(len(levels) >= 3, f"code indentation preserved, levels {levels}")

        # ---- online table --------------------------------------------------
        on = doc.tables[4].rows[1:]
        check(len(on) == 1 and on[0].cells[1].text.strip() == "(A)"
              and on[0].cells[3].text.strip() == "(09)"
              and on[0].cells[4].text.strip() == "C",
              "online = ONE integrated Q-3(A), marks (09), Bloom C (§9)")
        body = on[0].cells[2].text
        check("algorithm" in body.lower() and "flowchart" in body.lower(),
              "online question requires an algorithm and a flowchart (Unit 1)")
        for probe in ("loop", "again", "how many", "Rs. 2000"):
            check(probe in body, f"online question requires {probe!r}")
        # the closing sentence forbids later-unit features, so judge the rest
        spec = body.split("Use of functions")[0]
        for banned in ("def ", "list(", "tuple(", "dict(", "set(", "import ",
                       "open(", "[", "lambda"):
            check(banned not in spec, f"online question avoids {banned!r} (§9)")
        check("ternary" not in body.lower(), "online question avoids ternary (§9)")

        # ---- answer key + confidentiality ----------------------------------
        block_txt = audit_txt.split(f"--- ANSWER KEY SET {letter}")[1] \
                             .split("--- ANSWER KEY")[0]
        key = dict(re.findall(r"(Q-1 \d+\))\s+(\([a-f]\))", block_txt))
        for sub, body, _ in mcq_rows:
            o = parse_options(body)
            m = by_opts[frozenset(o.values())]
            right = [L for L, v in o.items() if v == m["opts"][m["ans"]]]
            check(len(right) == 1 and key.get(f"Q-1 {sub}") == f"({right[0]})",
                  f"key Q-1 {sub} = ({right[0]}) is the printed position of the "
                  f"correct option (PB{m['pb']})")
        student = " ".join(c.text for t in doc.tables for r in t.rows
                           for c in r.cells)
        for leak in ("Practice Book", "PB Sr", "blueprint", "Answer",
                     "faculty", "audit"):
            check(leak not in student, f"student paper does not mention {leak!r} (§11)")
        check(not re.search(r"\bOR\b", student), "no OR-type options")
        mains = {r.cells[0].text.strip() for r in rows[1:]
                 if r.cells[0].text.strip()}
        mains |= {r.cells[0].text.strip() for r in on
                  if r.cells[0].text.strip()}
        check(mains == {"Q-1", "Q-2", "Q-3"},
              f"main questions = {sorted(mains)} (max 7 allowed)")

        # ---- template controls ---------------------------------------------
        floating = [i for i, t in enumerate(doc.tables)
                    if t._tbl.find(qn("w:tblPr")).find(qn("w:tblpPr")) is not None]
        check(not floating,
              f"no floating question table (would not split across a page): {floating}")
        heights = {tr.find(qn("w:trPr")).find(qn("w:trHeight")).get(qn("w:val"))
                   for tr in doc.tables[1]._tbl.findall(qn("w:tr"))
                   if tr.find(qn("w:trPr")) is not None
                   and tr.find(qn("w:trPr")).find(qn("w:trHeight")) is not None}
        check(heights == {"624", "677"},
              f"template row heights preserved (624 header / 677 question), "
              f"got {heights}")
        widths = [g.get(qn("w:w"))
                  for g in doc.tables[1]._tbl.find(qn("w:tblGrid"))
                  .findall(qn("w:gridCol"))]
        check(widths == ["642", "497", "8396", "810", "841"],
              f"template column widths preserved {widths}")
        check(len(doc.tables) == 6 and len(doc.tables[2].rows) == 3
              and len(doc.tables[5].rows) == 3,
              "Bloom taxonomy tables kept on both pages (§10)")
        check(doc.tables[0].rows[0].cells[0].text.strip() == "Enrollment No:",
              "enrollment-number row untouched (§10)")

    # ---- §3 blueprint ------------------------------------------------------
    on_u2, on_u3 = Q3[0]["unit_online"]
    total = {1: unit_offline[1], 2: unit_offline[2] + on_u2,
             3: unit_offline[3] + on_u3}
    check(unit_offline == OFFLINE_UNIT_TARGET,
          f"offline unit split {unit_offline[1]}/{unit_offline[2]}/"
          f"{unit_offline[3]} = target 2/6/8 (§3)")
    check(total == T1_UNIT_BLUEPRINT,
          f"T1 unit total {total[1]}/{total[2]}/{total[3]} = blueprint 2/10/13 (§3)")


def verify_sets_identical():
    print("\n== sets differ only in MCQ sequence (§4, §7) ==")
    opts_per_set, order_per_set, tail = {}, {}, {}
    for letter in ("A", "B", "C"):
        doc = Document(str(OUTDIR /
                           f"SET {letter}_T1_PYTHON-1_TEST PAPER_SEM I_MDP.docx"))
        rows = doc.tables[1].rows
        opts_per_set[letter] = [parse_options(r.cells[2].text)
                                for r in rows[2:12]]
        order_per_set[letter] = [tuple(sorted(o.values()))
                                 for o in opts_per_set[letter]]
        tail[letter] = ([r.cells[2].text for r in rows[12:]]
                        + [r.cells[2].text for r in doc.tables[4].rows[1:]])
    for x, y in (("A", "B"), ("B", "C"), ("A", "C")):
        check(sorted(order_per_set[x]) == sorted(order_per_set[y]),
              f"SET {x} and SET {y} contain the same ten questions")
        check(order_per_set[x] != order_per_set[y],
              f"SET {x} and SET {y} use a different MCQ sequence")
        check(opts_per_set[x] != opts_per_set[y],
              f"SET {x} and SET {y} place the questions differently")
    check(tail["A"] == tail["B"] == tail["C"],
          "Q-2 and Q-3 are byte-identical in all three sets")
    for letter in ("A", "B", "C"):
        letters = [o for opt in opts_per_set[letter]
                   for L, o in opt.items() if o ==
                   next(m["opts"][m["ans"]] for m in MCQS
                        if frozenset(m["opts"]) == frozenset(opt.values()))]
        from collections import Counter
        c = Counter(letters)
        check(max(c.values()) <= 3 and len(set(letters)) >= 4,
              f"SET {letter} key spread {dict(sorted(c.items()))}")


def main():
    verify_answers()
    verify_documents()
    verify_sets_identical()
    print()
    if FAIL:
        print(f"{len(FAIL)} CHECK(S) FAILED")
        for f in FAIL:
            print("   -", f)
        sys.exit(1)
    print("ALL CHECKS PASSED")


if __name__ == "__main__":
    main()
